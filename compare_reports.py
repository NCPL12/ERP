import openpyxl
from collections import defaultdict

DL = r"C:\Users\tekni\Downloads"
monthly_f = DL + r"\20260702063509monthly_stock_movement.xlsx"
dc_f = DL + r"\20260702063757dc_by_date.xlsx"
grn_f = DL + r"\20260702063758grn_by_date.xlsx"

def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0

# ---- Monthly report ----
wb = openpyxl.load_workbook(monthly_f, data_only=True)
ws = wb.active
# find sub-header row (contains "Qty") to detect layout
header_row = None
for r in range(1, 10):
    vals = [ws.cell(r, c).value for c in range(1, 17)]
    if "Qty" in [v for v in vals if isinstance(v, str)]:
        header_row = r
        sub = vals
        break
print("Monthly sub-header row:", header_row, "->", [v for v in sub if v not in (None, "")])

# detect old (combined Particulars, 14 cols) vs new (Model No + Description, 15 cols)
# labels for cols 1-3 live in the GROUP row (merged cells), one row above the sub-header
group_row = header_row - 1
new_layout = ws.cell(group_row, 3).value == "Description"
print("Monthly group row:", [ws.cell(group_row, c).value for c in range(1, 5)])
print("Monthly layout:", "NEW (Model No / Description)" if new_layout else "OLD (Particulars combined)")

monthly = {}
for r in range(header_row + 1, ws.max_row + 1):
    c0 = ws.cell(r, 1).value
    if c0 in (None, "Total"):
        label = ws.cell(r, 1).value
        if label == "Total":
            if new_layout:
                print("Monthly TOTAL row: inQty=%s inVal=%s outQty=%s outVal=%s" % (
                    ws.cell(r, 7).value, ws.cell(r, 9).value, ws.cell(r, 10).value, ws.cell(r, 12).value))
            else:
                print("Monthly TOTAL row: inQty=%s inVal=%s outQty=%s outVal=%s" % (
                    ws.cell(r, 6).value, ws.cell(r, 8).value, ws.cell(r, 9).value, ws.cell(r, 11).value))
        continue
    # skip rows whose first cell isn't a serial number (safety)
    if not isinstance(c0, (int, float)):
        continue
    if new_layout:
        model = str(ws.cell(r, 2).value or "").strip()
        base = 4  # openQty col
    else:
        part = str(ws.cell(r, 2).value or "")
        model = part.split(" / ")[0].strip()
        base = 3
    if not model:
        continue
    monthly[model] = {
        "openQty": num(ws.cell(r, base).value),
        "inQty":  num(ws.cell(r, base + 3).value),
        "inRate": num(ws.cell(r, base + 4).value),
        "inVal":  num(ws.cell(r, base + 5).value),
        "outQty": num(ws.cell(r, base + 6).value),
        "outRate":num(ws.cell(r, base + 7).value),
        "outVal": num(ws.cell(r, base + 8).value),
        "clQty":  num(ws.cell(r, base + 9).value),
    }
print("Monthly items:", len(monthly))

# ---- DC by date ----
wb = openpyxl.load_workbook(dc_f, data_only=True)
ws = wb.active
hdr = None
for r in range(1, 10):
    if ws.cell(r, 1).value == "DC No":
        hdr = r
        break
print("\nDC header row:", hdr, "cols:", [ws.cell(hdr, c).value for c in range(1, 8)])
dc = defaultdict(lambda: {"qty": 0.0, "val": 0.0, "rates": set(), "rows": 0})
for r in range(hdr + 1, ws.max_row + 1):
    model = str(ws.cell(r, 3).value or "").strip()
    if not model:
        continue
    d = dc[model]
    d["qty"] += num(ws.cell(r, 5).value)
    d["val"] += num(ws.cell(r, 7).value)
    d["rates"].add(round(num(ws.cell(r, 6).value), 2))
    d["rows"] += 1
print("DC models:", len(dc), "total qty:", sum(d['qty'] for d in dc.values()),
      "total val:", round(sum(d['val'] for d in dc.values()), 2))

# ---- GRN by date ----
wb = openpyxl.load_workbook(grn_f, data_only=True)
ws = wb.active
hdr = None
for r in range(1, 10):
    if ws.cell(r, 1).value == "Model No":
        hdr = r
        break
print("\nGRN header row:", hdr, "cols:", [ws.cell(hdr, c).value for c in range(1, 8)])
grn = defaultdict(lambda: {"qty": 0.0, "val": 0.0, "rates": set(), "rows": 0})
for r in range(hdr + 1, ws.max_row + 1):
    model = str(ws.cell(r, 1).value or "").strip()
    if not model:
        continue
    g = grn[model]
    g["qty"] += num(ws.cell(r, 3).value)
    g["val"] += num(ws.cell(r, 6).value)
    g["rates"].add(round(num(ws.cell(r, 5).value), 2))
    g["rows"] += 1
print("GRN models:", len(grn), "total qty:", sum(g['qty'] for g in grn.values()),
      "total val:", round(sum(g['val'] for g in grn.values()), 2))

TOL = 0.05

print("\n==== GRN vs Monthly INWARD ====")
mismatch = 0
for model in sorted(set(list(grn.keys()) + [m for m in monthly if monthly[m]["inQty"]])):
    g = grn.get(model)
    m = monthly.get(model)
    gq = g["qty"] if g else 0.0
    gv = g["val"] if g else 0.0
    mq = m["inQty"] if m else 0.0
    mv = m["inVal"] if m else 0.0
    if abs(gq - mq) > TOL or abs(gv - mv) > TOL:
        mismatch += 1
        grates = sorted(g["rates"]) if g else []
        mr = m["inRate"] if m else 0.0
        print(f"{model}: GRN qty={gq} val={round(gv,2)} rates={grates} | Monthly inQty={mq} inVal={round(mv,2)} rate={mr}")
if mismatch == 0:
    print("ALL MATCH")
else:
    print(f"-> {mismatch} model(s) mismatch")

print("\n==== DC vs Monthly OUTWARD ====")
mismatch = 0
for model in sorted(set(list(dc.keys()) + [m for m in monthly if monthly[m]["outQty"]])):
    d = dc.get(model)
    m = monthly.get(model)
    dq = d["qty"] if d else 0.0
    dv = d["val"] if d else 0.0
    mq = m["outQty"] if m else 0.0
    mv = m["outVal"] if m else 0.0
    if abs(dq - mq) > TOL or abs(dv - mv) > TOL:
        mismatch += 1
        drates = sorted(d["rates"]) if d else []
        mr = m["outRate"] if m else 0.0
        print(f"{model}: DC qty={dq} val={round(dv,2)} rates={drates} | Monthly outQty={mq} outVal={round(mv,2)} rate={mr}")
if mismatch == 0:
    print("ALL MATCH")
else:
    print(f"-> {mismatch} model(s) mismatch")
